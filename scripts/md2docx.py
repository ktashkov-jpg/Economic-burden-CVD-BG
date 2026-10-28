"""Manuscript .md -> .docx: drops the status line and Appendix A, fixes figure paths, styles tables.
usage: python md2docx.py in.md out.docx figdir_from figdir_to"""
import sys, re, subprocess, tempfile, os
src,out,ffrom,fto=sys.argv[1:5]
JOURNAL = len(sys.argv) > 5 and sys.argv[5] == "--journal"
s=open(src,encoding='utf-8',newline='').read().replace('\r\n','\n')
s=re.sub(r'^\*\*Manuscript status:\*\*.*\n','',s,flags=re.M)
if '## Appendix A' in s: s=s[:s.index('## Appendix A')].rstrip()
s=s.rstrip()
if s.endswith('---'): s=s[:-3].rstrip()
s+='\n'
if ffrom and ffrom!=fto and ('/' in ffrom or '\\' in ffrom): s=s.replace(ffrom,fto)
# Unicode superscript citations -> real superscript runs (Times New Roman lacks most of these glyphs)
SUP=str.maketrans('⁰¹²³⁴⁵⁶⁷⁸⁹⁻','0123456789–')
s=re.sub('[⁰¹²³⁴-⁹⁻]+',lambda m:'^'+m.group(0).translate(SUP)+'^',s)
s=s.replace('^,^',',')
lines=s.split('\n'); title=lines[0][2:]; rest=lines[1:]
i=rest.index('---') if '---' in rest else 0
head=[l for l in rest[:i] if l.strip()]
s='\n\n'.join(head)+'\n\n'+'\n'.join(rest[i:])
s=re.sub(r'^\|(?: *:?-+:? *\|)+$', lambda m: '|'+'---|'*(m.group(0).count('|')-1), s, flags=re.M)
s='---\ntitle: "'+title.replace('"','\\"')+'"\n---\n\n'+s
d=tempfile.mkdtemp(); md=os.path.join(d,'e.md'); raw=os.path.join(d,'raw.docx')
open(md,'w',encoding='utf-8').write(s)
subprocess.run(['pandoc',md,'-o',raw],check=True)
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
doc=Document(raw)
STY={st.element.get(qn('w:styleId')): st for st in doc.styles}
class _S(dict):
    def __getitem__(self, k): return STY[k.replace(' ','')]
for name in ['Normal','Body Text','First Paragraph','Compact','Title','Heading 1','Heading 2','Heading 3']:
    try:
        st=_S()[name]; st.font.name='Calibri'
        rf=st.element.get_or_add_rPr().find(qn('w:rFonts'))
        for a in ['w:ascii','w:hAnsi','w:eastAsia','w:cs']: rf.set(qn(a),'Calibri')
        for a in ['w:asciiTheme','w:hAnsiTheme','w:eastAsiaTheme','w:cstheme']:
            if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
        if name in ['Normal','Body Text','First Paragraph']: st.font.size=Pt(11)
    except KeyError: pass
_S()['Title'].font.size=Pt(16)
for sec in doc.sections: sec.left_margin=sec.right_margin=Cm(2.2)
TOTAL=Cm(16.6)
for t in doc.tables:
    tblPr=t._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for e in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+e); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'999999'); b.append(el)
    tblPr.append(b)
    lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
    ncol=len(t.columns)
    L=[max(11,max(min(len(r.cells[c].text),30 if c==0 else 22) for r in t.rows)) for c in range(ncol)]
    W=[int(TOTAL*l/sum(L)) for l in L]
    tw=tblPr.find(qn('w:tblW'))
    if tw is not None: tw.set(qn('w:type'),'dxa'); tw.set(qn('w:w'),str(sum(W)//635))
    for gc,w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')),W): gc.set(qn('w:w'),str(w//635))
    for row in t.rows:
        for c,cell in enumerate(row.cells):
            cell.width=W[c]
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
                for r in p.runs: r.font.size=Pt(9)
for sh in doc.inline_shapes:
    ratio=sh.height/sh.width; sh.width=Cm(16); sh.height=int(Cm(16)*ratio)
if JOURNAL:
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    for name in ['Normal','Body Text','First Paragraph','Compact','Title','Heading 1','Heading 2','Heading 3']:
        try:
            st=_S()[name]; st.font.name='Times New Roman'
            from docx.shared import RGBColor; st.font.color.rgb=RGBColor(0,0,0)
            c=st.element.get_or_add_rPr().find(qn('w:color'))
            if c is not None:
                for a in ['w:themeColor','w:themeShade','w:themeTint']:
                    if c.get(qn(a)) is not None: del c.attrib[qn(a)]
            rf=st.element.get_or_add_rPr().find(qn('w:rFonts'))
            for a in ['w:ascii','w:hAnsi','w:eastAsia','w:cs']: rf.set(qn(a),'Times New Roman')
            st.font.size=Pt(12)
            if name not in ('Compact',): st.paragraph_format.line_spacing=2.0
        except KeyError: pass
    for name,b,i in [('Heading 1',True,False),('Heading 2',True,False),('Heading 3',True,True),('Title',True,False)]:
        try: st=_S()[name]; st.font.bold=b; st.font.italic=i; st.font.size=Pt(12)
        except KeyError: pass
    for sec in doc.sections:
        sec.page_width=Inches(8.5); sec.page_height=Inches(11)
        sec.left_margin=sec.right_margin=sec.top_margin=sec.bottom_margin=Inches(1)
        p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(); 
        for t,v in [('begin',None),(None,'PAGE'),('end',None)]:
            if t: e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),t)
            else: e=OxmlElement('w:instrText'); e.set(qn('xml:space'),'preserve'); e.text=v
            r._r.append(e)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing=1.0
                    for r in p.runs: r.font.size=Pt(10); r.font.name='Times New Roman'
doc.save(out)
print('wrote',out)
